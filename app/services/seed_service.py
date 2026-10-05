import os
import openpyxl
from datetime import datetime
from app.models import db
from app.models.user import User, PatientProfile
from app.models.doctor import Doctor, Specialization
from app.models.medicine import Medicine, MedicineCategory
from app.models.order import Order, OrderItem
from app.models.appointment import Appointment
from app.models.ambulance import EmergencyRequest

def migrate_and_seed_data(base_dir=None):
    """
    Reads existing Excel files from base_dir and migrates them into PostgreSQL / SQLAlchemy DB.
    Also creates initial admin, doctor, and patient demo accounts.
    """
    if base_dir is None:
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))

    print("[Seed Service] Starting data migration from Excel files...")

    # 1. Seed Default Users & Roles if not exists
    admin_user = User.query.filter_by(email='admin@reachme.com').first()
    if not admin_user:
        admin_user = User(
            email='admin@reachme.com',
            full_name='System Admin',
            phone='+1 800 555 0199',
            role='ADMIN'
        )
        admin_user.set_password('admin123')
        db.session.add(admin_user)

    demo_patient = User.query.filter_by(email='patient@reachme.com').first()
    if not demo_patient:
        demo_patient = User(
            email='patient@reachme.com',
            full_name='Chirayu Sahu',
            phone='+91 72489 46823',
            role='PATIENT'
        )
        demo_patient.set_password('patient123')
        db.session.add(demo_patient)
        db.session.flush()
        
        patient_profile = PatientProfile(
            user_id=demo_patient.id,
            blood_group='O+',
            address='VIT Vellore, Tamil Nadu',
            emergency_contact='+91 98765 43210'
        )
        db.session.add(patient_profile)

    demo_doctor = User.query.filter_by(email='doctor@reachme.com').first()
    if not demo_doctor:
        demo_doctor = User(
            email='doctor@reachme.com',
            full_name='Dr. John Smith',
            phone='+91 99000 11223',
            role='DOCTOR'
        )
        demo_doctor.set_password('doctor123')
        db.session.add(demo_doctor)

    dr_johnsmith = User.query.filter_by(email='dr.johnsmith@reachme.com').first()
    if not dr_johnsmith:
        dr_johnsmith = User(
            email='dr.johnsmith@reachme.com',
            full_name='Dr. John Smith',
            phone='+91 99000 11223',
            role='DOCTOR'
        )
        dr_johnsmith.set_password('doctor123')
        db.session.add(dr_johnsmith)

    db.session.commit()

    # 2. Migrate Medicine Prices (medicine_prices.xlsx)
    med_file = os.path.join(base_dir, 'medicine_prices.xlsx')
    if os.path.exists(med_file):
        try:
            wb = openpyxl.load_workbook(med_file)
            ws = wb.active
            gen_cat = MedicineCategory.query.filter_by(name='General Care').first()
            if not gen_cat:
                gen_cat = MedicineCategory(name='General Care', description='Over-the-counter and general healthcare medicines')
                db.session.add(gen_cat)
                db.session.commit()

            antibiotic_cat = MedicineCategory.query.filter_by(name='Antibiotics').first()
            if not antibiotic_cat:
                antibiotic_cat = MedicineCategory(name='Antibiotics', description='Prescription antibiotics')
                db.session.add(antibiotic_cat)
                db.session.commit()

            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                med_name = str(row[0]).strip()
                price = float(row[1]) if row[1] is not None else 50.0
                
                existing = Medicine.query.filter(Medicine.name.ilike(med_name)).first()
                if not existing:
                    # Categorization logic
                    is_antibiotic = any(k in med_name.lower() for k in ['cillin', 'mycin', 'cycline', 'dazole', 'floxacin'])
                    cat_id = antibiotic_cat.id if is_antibiotic else gen_cat.id
                    requires_rx = is_antibiotic or any(k in med_name.lower() for k in ['zepam', 'pram', 'prazole'])
                    
                    med = Medicine(
                        name=med_name,
                        price=price,
                        category_id=cat_id,
                        stock_quantity=150,
                        requires_prescription=requires_rx,
                        is_available=True
                    )
                    db.session.add(med)
                    count += 1
            db.session.commit()
            print(f"[Seed Service] Migrated {count} medicines from Excel.")
        except Exception as e:
            print(f"[Seed Service Error] Failed to migrate medicines: {e}")

    # 3. Migrate Doctors (doctor_list.xlsx)
    doc_file = os.path.join(base_dir, 'doctor_list.xlsx')
    if os.path.exists(doc_file):
        try:
            wb = openpyxl.load_workbook(doc_file)
            ws = wb.active
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                name = str(row[0]).strip()
                spec_name = str(row[1]).strip() if len(row) > 1 and row[1] else 'General Physician'
                
                # Specialization ORM
                spec_obj = Specialization.query.filter_by(name=spec_name).first()
                if not spec_obj:
                    spec_obj = Specialization(name=spec_name, description=f'Specialized medical care in {spec_name}')
                    db.session.add(spec_obj)
                    db.session.flush()

                existing_doc = Doctor.query.filter_by(name=name).first()
                if not existing_doc:
                    # Create a user login account for doctor
                    clean_email = name.lower().replace('dr.', '').replace(' ', '').strip() + '@reachme.com'
                    doc_user = User.query.filter_by(email=clean_email).first()
                    if not doc_user:
                        doc_user = User(
                            email=clean_email,
                            full_name=name,
                            phone='+91 99000 11223',
                            role='DOCTOR'
                        )
                        doc_user.set_password('doctor123')
                        db.session.add(doc_user)
                        db.session.flush()

                    doc = Doctor(
                        user_id=doc_user.id,
                        name=name,
                        specialty=spec_name,
                        specialization_id=spec_obj.id,
                        experience_years=8 + (count % 15),
                        consultation_fee=400.0 + ((count % 5) * 100),
                        rating=4.5 + ((count % 5) * 0.1),
                        bio=f'{name} is a senior {spec_name} with over {8 + (count % 15)} years of clinical practice.',
                        is_available=True
                    )
                    db.session.add(doc)
                    count += 1
            db.session.commit()
            print(f"[Seed Service] Migrated {count} doctors from Excel.")
        except Exception as e:
            print(f"[Seed Service Error] Failed to migrate doctors: {e}")

    # 4. Migrate Consultations (consultations.xlsx)
    cons_file = os.path.join(base_dir, 'consultations.xlsx')
    if os.path.exists(cons_file):
        try:
            wb = openpyxl.load_workbook(cons_file)
            ws = wb.active
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                p_name = str(row[0]).strip()
                p_phone = str(row[1]).strip() if len(row) > 1 and row[1] else '+91 0000000000'
                problem = str(row[2]).strip() if len(row) > 2 and row[2] else 'General Consultation'
                doc_info = str(row[3]).strip() if len(row) > 3 and row[3] else 'Dr. John Smith'
                date_str = str(row[4]).strip() if len(row) > 4 and row[4] else '2026-10-10'
                time_slot = str(row[5]).strip() if len(row) > 5 and row[5] else '10:00 AM'

                doc_name = doc_info.split(':')[0].strip() if ':' in doc_info else doc_info
                doctor = Doctor.query.filter(Doctor.name.ilike(f"%{doc_name}%")).first()
                if not doctor:
                    doctor = Doctor.query.first()

                if doctor:
                    app = Appointment(
                        patient_name=p_name,
                        patient_phone=p_phone,
                        problem_description=problem,
                        doctor_id=doctor.id,
                        doctor_name=doctor.name,
                        doctor_specialty=doctor.specialty,
                        appointment_date=date_str,
                        time_slot=time_slot,
                        status='CONFIRMED'
                    )
                    db.session.add(app)
                    count += 1
            db.session.commit()
            print(f"[Seed Service] Migrated {count} consultations from Excel.")
        except Exception as e:
            print(f"[Seed Service Error] Failed to migrate consultations: {e}")

    # 5. Migrate Orders (orders.xlsx)
    ord_file = os.path.join(base_dir, 'orders.xlsx')
    if os.path.exists(ord_file):
        try:
            wb = openpyxl.load_workbook(ord_file)
            ws = wb.active
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                name = str(row[0]).strip()
                phone = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                med_name = str(row[2]).strip() if len(row) > 2 and row[2] else 'Paracetamol'
                qty = int(row[3]) if len(row) > 3 and row[3] else 1
                addr = str(row[4]).strip() if len(row) > 4 and row[4] else 'Address not specified'

                med = Medicine.query.filter(Medicine.name.ilike(med_name)).first()
                unit_price = med.price if med else 20.0
                med_id = med.id if med else 1
                actual_name = med.name if med else med_name

                order = Order(
                    full_name=name,
                    phone_number=phone,
                    shipping_address=addr,
                    total_amount=unit_price * qty,
                    status='DELIVERED'
                )
                db.session.add(order)
                db.session.flush()

                item = OrderItem(
                    order_id=order.id,
                    medicine_id=med_id,
                    medicine_name=actual_name,
                    quantity=qty,
                    unit_price=unit_price,
                    subtotal=unit_price * qty
                )
                db.session.add(item)
                count += 1
            db.session.commit()
            print(f"[Seed Service] Migrated {count} orders from Excel.")
        except Exception as e:
            print(f"[Seed Service Error] Failed to migrate orders: {e}")

    print("[Seed Service] Data migration complete!")
